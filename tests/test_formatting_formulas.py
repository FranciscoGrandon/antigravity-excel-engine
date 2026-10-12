# Creado por Francisco Grandón Vergara
"""
Pruebas unitarias para los submódulos de formato y fórmulas de antigravity_excel_engine.
"""

import unittest
from openpyxl import Workbook

from antigravity_excel_engine.formatting import (
    FinancialStyleGuide,
    hex_to_rgb_int,
    hex_to_argb_hex,
    apply_number_format,
    apply_style_preset_openpyxl,
    apply_style_preset,
)
from antigravity_excel_engine.formulas import (
    MODERN_EXCEL_FUNCTIONS,
    sanitize_formula_compatibility,
    strip_xlfn_prefix,
    is_formula_error_value,
    get_formula_error_name,
    audit_openpyxl_formula_errors,
    propagate_formula_row,
    replace_formula_placeholders,
    make_absolute_reference,
    make_relative_reference,
    make_row_absolute_reference,
    make_col_absolute_reference,
)


class TestFormattingSubmodule(unittest.TestCase):
    def setUp(self):
        self.wb = Workbook()
        self.ws = self.wb.active

    def test_financial_style_guide_constants(self):
        self.assertEqual(FinancialStyleGuide.CURRENCY_FORMAT, "$ #,##0.00;($ #,##0.00);\"-\"")
        self.assertEqual(FinancialStyleGuide.PERCENT_FORMAT, "0.0%")
        self.assertEqual(FinancialStyleGuide.COLOR_PRIMARY, "#0E2E63")
        self.assertIn("header", FinancialStyleGuide.PRESETS)
        self.assertIn("input_cell", FinancialStyleGuide.PRESETS)
        self.assertIn("total_row", FinancialStyleGuide.PRESETS)

    def test_hex_conversions(self):
        # #0E2E63 -> R=14, G=46, B=99 -> BGR int = 14 + (46 << 8) + (99 << 16)
        bgr = hex_to_rgb_int("#0E2E63")
        self.assertEqual(bgr, 14 + (46 << 8) + (99 << 16))
        self.assertEqual(hex_to_argb_hex("#0E2E63"), "0E2E63")
        self.assertEqual(hex_to_argb_hex("FFF2CC"), "FFF2CC")

    def test_apply_style_preset_openpyxl_header(self):
        cell = self.ws["A1"]
        cell.value = "Test Header"
        apply_style_preset_openpyxl([cell], "header")
        self.assertTrue(cell.font.bold)
        self.assertEqual(cell.font.color.rgb, "00FFFFFF")
        self.assertEqual(cell.fill.start_color.rgb, "000E2E63")

    def test_apply_style_preset_openpyxl_input_cell(self):
        cell = self.ws["B2"]
        apply_style_preset(cell, "input_cell")
        self.assertEqual(cell.fill.start_color.rgb, "00FFF2CC")
        self.assertEqual(cell.font.color.rgb, "001F4E79")
        self.assertEqual(cell.border.left.style, "medium")

    def test_apply_style_preset_openpyxl_total_row(self):
        cell = self.ws["C10"]
        apply_style_preset_openpyxl([cell], "total_row")
        self.assertTrue(cell.font.bold)
        self.assertEqual(cell.border.top.style, "thin")
        self.assertEqual(cell.border.bottom.style, "double")

    def test_apply_number_format(self):
        cell = self.ws["D5"]
        apply_number_format(cell, FinancialStyleGuide.CURRENCY_FORMAT)
        self.assertEqual(cell.number_format, FinancialStyleGuide.CURRENCY_FORMAT)

        cell2 = self.ws["E5"]
        apply_style_preset(cell2, "currency_clp")
        self.assertEqual(cell2.number_format, FinancialStyleGuide.CURRENCY_CLP)


class TestFormulasSubmodule(unittest.TestCase):
    def setUp(self):
        self.wb = Workbook()
        self.ws = self.wb.active

    def test_sanitize_modern_functions(self):
        # Text/Logic
        f1 = sanitize_formula_compatibility('=TEXTJOIN(",", TRUE, A1:A5)')
        self.assertEqual(f1, '=_xlfn.TEXTJOIN(",", TRUE, A1:A5)')

        # Dynamic arrays (Excel 365)
        f2 = sanitize_formula_compatibility('=XLOOKUP(A1, B:B, C:C)')
        self.assertEqual(f2, '=_xlfn.XLOOKUP(A1, B:B, C:C)')

        f3 = sanitize_formula_compatibility('=UNIQUE(FILTER(A1:A10, B1:B10>0))')
        self.assertEqual(f3, '=_xlfn.UNIQUE(_xlfn.FILTER(A1:A10, B1:B10>0))')

        f4 = sanitize_formula_compatibility('=LET(x, 5, x*2)')
        self.assertEqual(f4, '=_xlfn.LET(x, 5, x*2)')

        # Strip xlfn
        cleaned = strip_xlfn_prefix(f3)
        self.assertEqual(cleaned, '=UNIQUE(FILTER(A1:A10, B1:B10>0))')

    def test_formula_placeholders_and_propagation(self):
        tmpl = "=A{r} * $B$1 + {tax_rate}"
        replaced = replace_formula_placeholders(tmpl, row=5, extra_vars={"tax_rate": 0.19})
        self.assertEqual(replaced, "=A5 * $B$1 + 0.19")

        # propagate_formula_row
        prop = propagate_formula_row("=A2 + $C$4 + B2", row_idx=10, base_row=2)
        self.assertEqual(prop, "=A10 + $C$4 + B10")

    def test_reference_helpers(self):
        self.assertEqual(make_absolute_reference("C", 4), "$C$4")
        self.assertEqual(make_relative_reference("b", 12), "B12")
        self.assertEqual(make_row_absolute_reference("a", 2), "A$2")
        self.assertEqual(make_col_absolute_reference("f", 7), "$F7")

    def test_formula_error_audit(self):
        self.assertTrue(is_formula_error_value("#VALUE!"))
        self.assertTrue(is_formula_error_value("#REF!"))
        self.assertTrue(is_formula_error_value("#DIV/0!"))
        self.assertTrue(is_formula_error_value(-2146826281)) # COM DIV/0
        self.assertFalse(is_formula_error_value(123))
        self.assertFalse(is_formula_error_value("Valid Text"))

        self.assertEqual(get_formula_error_name(-2146826273), "#VALUE!")
        self.assertEqual(get_formula_error_name("#SPILL!"), "#SPILL!")

        # Audit openpyxl worksheet
        self.ws["A1"] = "OK"
        self.ws["A2"] = "#DIV/0!"
        self.ws["B2"] = "#N/A"
        errors = audit_openpyxl_formula_errors(self.ws, "A1:B2")
        self.assertEqual(len(errors), 2)
        coords = [e["cell"] for e in errors]
        self.assertIn("A2", coords)
        self.assertIn("B2", coords)


if __name__ == "__main__":
    unittest.main()
