# Creado por Francisco Grandón Vergara
"""
Funciones universales para aplicar estilos, formatos de celda y presets en openpyxl y COM.
"""

from typing import Any, Dict, List, Optional, Union

from .presets import FinancialStyleGuide


def hex_to_rgb_int(hex_str: str) -> int:
    """Convierte '#RRGGBB' o 'RRGGBB' a entero BGR para Excel COM (0x00BBGGRR)."""
    clean_hex = str(hex_str).lstrip('#')
    if len(clean_hex) == 6:
        r = int(clean_hex[0:2], 16)
        g = int(clean_hex[2:4], 16)
        b = int(clean_hex[4:6], 16)
        return r + (g << 8) + (b << 16)
    return 0


def hex_to_argb_hex(hex_str: str) -> str:
    """Normaliza un color hex a formato ARGB o RRGGBB para openpyxl."""
    clean = str(hex_str).lstrip('#').upper()
    if len(clean) == 6:
        return clean
    elif len(clean) == 8:
        return clean
    return "000000"


def apply_number_format(target: Any, fmt: str, is_com: bool = False) -> None:
    """
    Aplica una máscara de formato numérico a una celda openpyxl o rango COM.
    """
    if not fmt:
        return
    if is_com:
        FinancialStyleGuide.apply_number_format_resilient(target, fmt)
    else:
        # openpyxl cell or list of cells
        if hasattr(target, "number_format"):
            target.number_format = fmt
        elif hasattr(target, "__iter__"):
            for c in target:
                if hasattr(c, "number_format"):
                    c.number_format = fmt


def apply_borders_openpyxl(
    cells: List[Any],
    top_style: Optional[str] = None,
    top_color: str = "000000",
    bottom_style: Optional[str] = None,
    bottom_color: str = "000000",
    left_style: Optional[str] = None,
    left_color: str = "000000",
    right_style: Optional[str] = None,
    right_color: str = "000000",
) -> None:
    """Aplica bordes específicos a una lista de celdas de openpyxl preservando lados intactos."""
    from openpyxl.styles import Border, Side

    for cell in cells:
        cur_border = getattr(cell, "border", None)
        top_side = Side(style=top_style, color=hex_to_argb_hex(top_color)) if top_style else (cur_border.top if cur_border else None)
        bottom_side = Side(style=bottom_style, color=hex_to_argb_hex(bottom_color)) if bottom_style else (cur_border.bottom if cur_border else None)
        left_side = Side(style=left_style, color=hex_to_argb_hex(left_color)) if left_style else (cur_border.left if cur_border else None)
        right_side = Side(style=right_style, color=hex_to_argb_hex(right_color)) if right_style else (cur_border.right if cur_border else None)

        cell.border = Border(top=top_side, bottom=bottom_side, left=left_side, right=right_side)


def apply_style_preset_openpyxl(cells: List[Any], preset: str) -> None:
    """
    Aplica un preset de diseño estándar a una colección de celdas de openpyxl.
    Presets admitidos:
      - header
      - input_cell
      - total_row
      - kpi_card
      - zebra_even
      - zebra_odd
      - delta_positive
      - delta_negative
      - currency
      - currency_clp
      - percent
      - date
    """
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    p = preset.lower()
    if p == "header":
        fill = PatternFill(start_color="0E2E63", end_color="0E2E63", fill_type="solid")
        for cell in cells:
            cell.fill = fill
            font_name = cell.font.name if cell.font and cell.font.name else "Segoe UI"
            font_size = cell.font.size if cell.font and cell.font.size else 11
            cell.font = Font(name=font_name, size=font_size, bold=True, color="FFFFFF")
            cell.alignment = Alignment(
                vertical="center",
                horizontal=cell.alignment.horizontal if cell.alignment else None,
                wrap_text=cell.alignment.wrap_text if cell.alignment else None
            )
    elif p == "input_cell":
        fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        border = Border(
            left=Side(style='medium', color='1F4E79'),
            right=Side(style='medium', color='1F4E79'),
            top=Side(style='medium', color='1F4E79'),
            bottom=Side(style='medium', color='1F4E79')
        )
        for cell in cells:
            cell.fill = fill
            font_name = cell.font.name if cell.font and cell.font.name else "Segoe UI"
            font_size = cell.font.size if cell.font and cell.font.size else 11
            bold = cell.font.bold if cell.font else False
            cell.font = Font(name=font_name, size=font_size, bold=bold, color="1F4E79")
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical=cell.alignment.vertical if cell.alignment else "center")
    elif p == "total_row":
        for cell in cells:
            font_name = cell.font.name if cell.font and cell.font.name else "Segoe UI"
            font_size = cell.font.size if cell.font and cell.font.size else 10
            cell.font = Font(name=font_name, size=font_size, bold=True, color=cell.font.color if cell.font else None)
            cell.border = Border(
                top=Side(style='thin', color='000000'),
                bottom=Side(style='double', color='000000'),
                left=cell.border.left if cell.border else None,
                right=cell.border.right if cell.border else None
            )
    elif p == "kpi_card":
        fill = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
        border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        for cell in cells:
            cell.fill = fill
            font_name = cell.font.name if cell.font and cell.font.name else "Segoe UI"
            font_size = cell.font.size if cell.font and cell.font.size else 12
            cell.font = Font(name=font_name, size=font_size, bold=True, color="0E2E63")
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
    elif p == "zebra_even":
        fill = PatternFill(start_color="F2F4F8", end_color="F2F4F8", fill_type="solid")
        for cell in cells:
            cell.fill = fill
    elif p == "zebra_odd":
        fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        for cell in cells:
            cell.fill = fill
    elif p == "currency":
        for cell in cells:
            cell.number_format = FinancialStyleGuide.CURRENCY_FORMAT
    elif p == "currency_clp":
        for cell in cells:
            cell.number_format = FinancialStyleGuide.CURRENCY_CLP
    elif p in ("percent", "percentage"):
        for cell in cells:
            cell.number_format = FinancialStyleGuide.PERCENT_FORMAT
    elif p == "date":
        for cell in cells:
            cell.number_format = FinancialStyleGuide.DATE
    elif p == "delta_positive":
        for cell in cells:
            font_name = cell.font.name if cell.font and cell.font.name else "Segoe UI"
            font_size = cell.font.size if cell.font and cell.font.size else 10
            cell.font = Font(name=font_name, size=font_size, bold=True, color="008000")
    elif p == "delta_negative":
        for cell in cells:
            font_name = cell.font.name if cell.font and cell.font.name else "Segoe UI"
            font_size = cell.font.size if cell.font and cell.font.size else 10
            cell.font = Font(name=font_name, size=font_size, bold=True, color="C00000")
    else:
        valid_presets = "header, input_cell, total_row, kpi_card, zebra_even, zebra_odd, currency, currency_clp, percent, date, delta_positive, delta_negative"
        raise ValueError(f"Preset no reconocido: '{preset}'. Presets válidos: {valid_presets}")


def apply_style_preset_com(rng: Any, preset: str) -> None:
    """
    Aplica un preset de diseño estándar a un objeto Range de win32com.
    """
    p = preset.lower()
    if p == "header":
        rng.Interior.Color = hex_to_rgb_int("#0E2E63")
        rng.Font.Color = hex_to_rgb_int("#FFFFFF")
        rng.Font.Bold = True
        rng.VerticalAlignment = -4108  # xlCenter
    elif p == "input_cell":
        rng.Interior.Color = hex_to_rgb_int("#FFF2CC")
        rng.Font.Color = hex_to_rgb_int("#1F4E79")
        rng.HorizontalAlignment = -4108  # xlCenter
        rng.Borders.LineStyle = 1        # xlContinuous
        rng.Borders.Weight = -4138       # xlMedium
        rng.Borders.Color = hex_to_rgb_int("#1F4E79")
    elif p == "total_row":
        rng.Font.Bold = True
        rng.Borders(8).LineStyle = 1     # xlEdgeTop = 8, xlContinuous = 1
        rng.Borders(8).Weight = 2        # xlThin = 2
        rng.Borders(9).LineStyle = -4119 # xlEdgeBottom = 9, xlDouble = -4119
    elif p == "kpi_card":
        rng.Interior.Color = hex_to_rgb_int("#F8F9FA")
        rng.Font.Color = hex_to_rgb_int("#0E2E63")
        rng.Font.Bold = True
        rng.Borders.LineStyle = 1        # xlContinuous
        rng.Borders.Weight = 2           # xlThin
        rng.Borders.Color = hex_to_rgb_int("#D9D9D9")
        rng.HorizontalAlignment = -4108
        rng.VerticalAlignment = -4108
    elif p == "zebra_even":
        rng.Interior.Color = hex_to_rgb_int("#F2F4F8")
    elif p == "zebra_odd":
        rng.Interior.Color = hex_to_rgb_int("#FFFFFF")
    elif p == "currency":
        FinancialStyleGuide.apply_number_format_resilient(rng, "currency")
    elif p == "currency_clp":
        FinancialStyleGuide.apply_number_format_resilient(rng, "currency_clp")
    elif p in ("percent", "percentage"):
        FinancialStyleGuide.apply_number_format_resilient(rng, "percent")
    elif p == "date":
        FinancialStyleGuide.apply_number_format_resilient(rng, "date")
    elif p == "delta_positive":
        rng.Font.Bold = True
        rng.Font.Color = hex_to_rgb_int("#008000")
    elif p == "delta_negative":
        rng.Font.Bold = True
        rng.Font.Color = hex_to_rgb_int("#C00000")
    else:
        valid_presets = "header, input_cell, total_row, kpi_card, zebra_even, zebra_odd, currency, currency_clp, percent, date, delta_positive, delta_negative"
        raise ValueError(f"Preset no reconocido: '{preset}'. Presets válidos: {valid_presets}")


def apply_style_preset(target: Any, preset: str, is_com: bool = False) -> None:
    """
    Punto de entrada universal para aplicar un preset de diseño tanto a objetos openpyxl como win32com Range.
    """
    if is_com:
        apply_style_preset_com(target, preset)
    else:
        if isinstance(target, list):
            cells = target
        elif hasattr(target, "__iter__") and not hasattr(target, "value"):
            cells = list(target)
        else:
            cells = [target]
        apply_style_preset_openpyxl(cells, preset)
