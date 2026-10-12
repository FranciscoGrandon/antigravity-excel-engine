# Creado por Francisco Grandón Vergara
"""
Backend Headless OpenPyXL de Antigravity Excel Engine.
Manipulación fuera de pantalla y en disco de archivos .xlsx sin requerir Microsoft Excel instalado.
"""

import os
import re
import io
import csv
import time
import datetime
import statistics
from typing import List, Dict, Any, Optional, Union

# Lazy Loaders para openpyxl, estilos, comentarios y módulos secundarios
_openpyxl_module = None
_openpyxl_styles = None
_openpyxl_comments = None

Font = None
PatternFill = None
Border = None
Side = None
Alignment = None
Comment = None


def _get_openpyxl():
    """Carga perezosa del módulo openpyxl con cacheo singleton."""
    global _openpyxl_module
    if _openpyxl_module is None:
        try:
            import openpyxl as _op
            _openpyxl_module = _op
        except ImportError:
            _openpyxl_module = False
    return _openpyxl_module if _openpyxl_module is not False else None


def _is_openpyxl_available() -> bool:
    """Verifica si openpyxl está instalado sin forzar su import si aún no ha sido cargado."""
    global _openpyxl_module
    if _openpyxl_module is not None:
        return _openpyxl_module is not False
    try:
        import importlib.util
        return importlib.util.find_spec("openpyxl") is not None
    except Exception:
        return _get_openpyxl() is not None


def _get_styles():
    global _openpyxl_styles, Font, PatternFill, Border, Side, Alignment
    if _openpyxl_styles is None:
        import openpyxl.styles as _st
        _openpyxl_styles = _st
        Font = _st.Font
        PatternFill = _st.PatternFill
        Border = _st.Border
        Side = _st.Side
        Alignment = _st.Alignment
    return _openpyxl_styles


def _get_comments():
    global _openpyxl_comments, Comment
    if _openpyxl_comments is None:
        import openpyxl.comments as _cm
        _openpyxl_comments = _cm
        Comment = _cm.Comment
    return _openpyxl_comments


win32file = None
win32con = None


def _get_win32file():
    global win32file
    if win32file is None:
        try:
            import win32file as _wfile
            win32file = _wfile
        except ImportError:
            pass
    return win32file


def _get_win32con():
    global win32con
    if win32con is None:
        try:
            import win32con as _wcon
            win32con = _wcon
        except ImportError:
            pass
    return win32con


from .base import BaseExcelBackend
from antigravity_excel_engine.utils.ranges import (
    coordinate_from_string,
    column_index_from_string,
    get_column_letter,
    range_boundaries,
)
from antigravity_excel_engine.utils.dates import parse_excel_date_serial
from antigravity_excel_engine.formatting import (
    FinancialStyleGuide,
    apply_style_preset_openpyxl,
)
from antigravity_excel_engine.formulas import (
    sanitize_formula_compatibility,
    audit_openpyxl_formula_errors,
    propagate_formula_row,
)

class HeadlessOpenPyXLBackend(BaseExcelBackend):
    def __init__(self, file_path: Optional[str] = None):
        op = _get_openpyxl()
        if not op:
            raise ImportError("openpyxl no está instalado en este entorno.")
        # Carga perezosa de estilos y comentarios en el namespace global
        _get_styles()
        _get_comments()
        self.file_path = file_path
        if file_path and os.path.exists(file_path):
            self.wb = self._load_workbook_resilient(file_path)
        else:
            self.wb = op.Workbook()

    def _load_workbook_resilient(self, file_path: str):
        """Carga openpyxl soportando lectura concurrente si el archivo está bloqueado por Excel."""
        op = _get_openpyxl()
        try:
            return op.load_workbook(file_path, data_only=False)
        except PermissionError:
            wfile = _get_win32file()
            wcon = _get_win32con()
            if wfile and wcon:
                handle = wfile.CreateFile(
                    os.path.abspath(file_path),
                    wfile.GENERIC_READ,
                    wfile.FILE_SHARE_READ | wfile.FILE_SHARE_WRITE | wfile.FILE_SHARE_DELETE,
                    None,
                    wfile.OPEN_EXISTING,
                    wfile.FILE_ATTRIBUTE_NORMAL,
                    None
                )
                chunks = []
                while True:
                    hr, data = wfile.ReadFile(handle, 64 * 1024)
                    if not data:
                        break
                    chunks.append(data)
                wfile.CloseHandle(handle)
                buffer = io.BytesIO(b"".join(chunks))
                return op.load_workbook(buffer, data_only=False)
            raise

    def _get_sheet(self, sheet_name: Optional[str]):
        if sheet_name:
            if sheet_name in self.wb.sheetnames:
                return self.wb[sheet_name]
            return self.wb.create_sheet(sheet_name)
        return self.wb.active

    def status(self) -> Dict[str, Any]:
        return {
            "backend": "headless_openpyxl",
            "file": self.file_path,
            "active_sheet": self.wb.active.title if self.wb.active else None,
            "sheets": self.wb.sheetnames
        }

    def get_cell_ranges(self, ranges: List[str], sheet: Optional[str] = None, include_formulas: bool = False, include_styles: bool = False) -> Dict[str, Any]:
        ws = self._get_sheet(sheet)
        res = {}
        for r in ranges:
            target_ws = ws
            clean_range = r
            if "!" in r:
                s_name, clean_range = r.split("!", 1)
                target_ws = self._get_sheet(s_name)

            min_col, min_row, max_col, max_row = range_boundaries(clean_range)
            vals = []
            forms = []
            for row_idx in range(min_row, max_row + 1):
                r_vals = []
                r_forms = []
                for col_idx in range(min_col, max_col + 1):
                    cell = target_ws.cell(row=row_idx, column=col_idx)
                    r_vals.append(cell.value)
                    r_forms.append(cell.value if (isinstance(cell.value, str) and cell.value.startswith("=")) else "")
                vals.append(r_vals)
                forms.append(r_forms)

            res[r] = {"values": vals}
            if include_formulas:
                res[r]["formulas"] = forms
        return res

    def set_cell_range(self, start_cell: str, values: List[List[Any]], sheet: Optional[str] = None, formulas: Optional[List[List[str]]] = None, styles: Optional[Dict[str, Any]] = None, autofit: bool = False) -> Dict[str, Any]:
        ws = self._get_sheet(sheet)
        start_col, start_row = coordinate_from_string(start_cell)
        start_col_idx = column_index_from_string(start_col)

        data = formulas if formulas else values
        if not data:
            return {"status": "error", "message": "No data provided"}

        for r_idx, row in enumerate(data):
            for c_idx, val in enumerate(row):
                cell = ws.cell(row=start_row + r_idx, column=start_col_idx + c_idx)
                if isinstance(val, str) and val.startswith("="):
                    cell.value = sanitize_formula_compatibility(val)
                else:
                    cell.value = val

        if autofit:
            for col in ws.columns:
                max_len = max(len(str(c.value or '')) for c in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

        return {"status": "success", "rows": len(data), "cols": len(data[0]) if data else 0}

    def set_cells(self, cells_map: Dict[str, Dict[str, Any]], sheet: Optional[str] = None, autofit: bool = False) -> Dict[str, Any]:
        ws = self._get_sheet(sheet)
        for coord, cell_info in cells_map.items():
            cell = ws[coord]
            if "formula" in cell_info and cell_info["formula"]:
                cell.value = sanitize_formula_compatibility(cell_info["formula"])
            elif "value" in cell_info:
                val = cell_info["value"]
                if isinstance(val, str) and val.startswith("="):
                    cell.value = sanitize_formula_compatibility(val)
                else:
                    cell.value = val

            if "note" in cell_info and cell_info["note"]:
                cell.comment = Comment(cell_info["note"], "Antigravity")

            styles = cell_info.get("cellStyles", {})
            if styles:
                font_args = {}
                if "fontWeight" in styles:
                    font_args["bold"] = (styles["fontWeight"] == "bold")
                if "fontStyle" in styles:
                    font_args["italic"] = (styles["fontStyle"] == "italic")
                if "fontSize" in styles:
                    font_args["size"] = int(styles["fontSize"])
                if "fontColor" in styles:
                    font_args["color"] = styles["fontColor"].lstrip("#")
                if font_args:
                    cell.font = Font(**font_args)

                if "backgroundColor" in styles:
                    hex_color = styles["backgroundColor"].lstrip("#")
                    cell.fill = PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

                if "numberFormat" in styles:
                    cell.number_format = styles["numberFormat"]

                if "horizontalAlignment" in styles:
                    cell.alignment = Alignment(horizontal=styles["horizontalAlignment"].lower())

        if autofit:
            for col in ws.columns:
                max_len = max(len(str(c.value or '')) for c in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

        return {"status": "success", "cells_updated": len(cells_map)}

    def get_range_as_csv(self, range_str: str, sheet: Optional[str] = None, delimiter: str = ";") -> str:
        data = self.get_cell_ranges([range_str], sheet=sheet)[range_str]["values"]
        output = io.StringIO()
        writer = csv.writer(output, delimiter=delimiter, lineterminator="\n")
        for row in data:
            writer.writerow(["" if x is None else str(x) for x in row])
        return output.getvalue().strip()

    def copy_paste_range(self, src_range: str, dst_start_cell: str, src_sheet: Optional[str] = None, dst_sheet: Optional[str] = None, paste_type: str = "all") -> Dict[str, Any]:
        ws_src = self._get_sheet(src_sheet)
        ws_dst = self._get_sheet(dst_sheet)

        min_col, min_row, max_col, max_row = range_boundaries(src_range)
        dst_col_letter, dst_row = coordinate_from_string(dst_start_cell)
        dst_col = column_index_from_string(dst_col_letter)

        for r_offset, r in enumerate(range(min_row, max_row + 1)):
            for c_offset, c in enumerate(range(min_col, max_col + 1)):
                src_cell = ws_src.cell(row=r, column=c)
                target_cell = ws_dst.cell(row=dst_row + r_offset, column=dst_col + c_offset)
                if paste_type in ("all", "values", "formulas"):
                    target_cell.value = src_cell.value
                if paste_type in ("all", "formats") and src_cell.has_style:
                    target_cell._style = src_cell._style

        return {"status": "success", "source": src_range, "destination": dst_start_cell}

    def check_formula_errors(self, range_str: str, sheet: Optional[str] = None) -> List[Dict[str, Any]]:
        ws = self._get_sheet(sheet)
        return audit_openpyxl_formula_errors(ws, range_str)

    def audit_data_quality(self, range_str: Optional[str] = None, sheet: Optional[str] = None) -> Dict[str, Any]:
        ws = self._get_sheet(sheet)
        if range_str:
            min_col, min_row, max_col, max_row = range_boundaries(range_str)
        else:
            min_col, min_row, max_col, max_row = 1, 1, ws.max_column, ws.max_row

        header_whitespace = []
        numbers_as_text = []
        inconsistent_types = []
        total_formulas = 0
        total_cells = (max_row - min_row + 1) * (max_col - min_col + 1)

        # Check headers (first row)
        for col_idx in range(min_col, max_col + 1):
            cell = ws.cell(row=min_row, column=col_idx)
            val = str(cell.value or "")
            if val and (val.startswith(" ") or val.endswith(" ") or "  " in val):
                header_whitespace.append({
                    "cell": cell.coordinate,
                    "header": val,
                    "trimmed": val.strip()
                })

        # Check column types, formulas, and numbers formatted as text
        for col_idx in range(min_col, max_col + 1):
            header_cell = ws.cell(row=min_row, column=col_idx)
            h_name = str(header_cell.value or f"Col_{col_idx}")
            col_types = set()
            for r_idx in range(min_row + 1, min(max_row + 1, min_row + 1000)):
                cell = ws.cell(row=r_idx, column=col_idx)
                val = cell.value
                if isinstance(val, str) and val.startswith("="):
                    total_formulas += 1
                if val is not None:
                    col_types.add(type(val).__name__)
                    if isinstance(val, str) and val.strip().replace(".", "", 1).replace("-", "", 1).isdigit():
                        if len(numbers_as_text) < 50:
                            numbers_as_text.append({
                                "cell": cell.coordinate,
                                "column": h_name,
                                "value": val
                            })
            if len(col_types) > 1:
                inconsistent_types.append({
                    "column_index": col_idx,
                    "header": h_name,
                    "types_detected": list(col_types)
                })

        range_name = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
        return {
            "sheet": ws.title,
            "range_audited": range_name,
            "total_cells": total_cells,
            "has_dynamic_formulas": (total_formulas > 0),
            "formula_count": total_formulas,
            "header_whitespace_issues": header_whitespace,
            "numbers_stored_as_text": numbers_as_text,
            "inconsistent_column_types": inconsistent_types
        }

    def aggregate(self, group_by_col: str, metric_col: str, agg_func: str = "sum", sheet: Optional[str] = None, top_n: int = 10, ascending: bool = False) -> Dict[str, Any]:
        ws = self._get_sheet(sheet)
        norm_group = group_by_col.strip().lower()
        norm_metric = metric_col.strip().lower()

        group_idx = None
        metric_idx = None

        for col_idx in range(1, ws.max_column + 1):
            h_val = str(ws.cell(row=1, column=col_idx).value or "").strip().lower()
            if h_val == norm_group and group_idx is None:
                group_idx = col_idx
            if h_val == norm_metric and metric_idx is None:
                metric_idx = col_idx

        if group_idx is None or metric_idx is None:
            raise ValueError(f"No se encontraron las columnas especificadas: group_by='{group_by_col}', metric='{metric_col}'")

        groups = {}
        for r_idx in range(2, ws.max_row + 1):
            g_val = ws.cell(row=r_idx, column=group_idx).value
            m_val = ws.cell(row=r_idx, column=metric_idx).value
            if g_val is None or m_val is None:
                continue

            try:
                num = float(m_val)
            except (ValueError, TypeError):
                continue

            g_key = str(g_val)
            cell_coord = ws.cell(row=r_idx, column=metric_idx).coordinate

            if g_key not in groups:
                groups[g_key] = {
                    "count": 0,
                    "values": [],
                    "cells": [],
                    "min_val": num,
                    "min_cell": cell_coord,
                    "max_val": num,
                    "max_cell": cell_coord
                }

            entry = groups[g_key]
            entry["count"] += 1
            entry["values"].append(num)
            entry["cells"].append(cell_coord)

            if num < entry["min_val"]:
                entry["min_val"] = num
                entry["min_cell"] = cell_coord
            if num > entry["max_val"]:
                entry["max_val"] = num
                entry["max_cell"] = cell_coord

        results = []
        for g_key, data in groups.items():
            vals = data["values"]
            total = sum(vals)
            avg = total / len(vals) if vals else 0.0

            if agg_func.lower() == "sum":
                agg_val = total
            elif agg_func.lower() in ("avg", "mean"):
                agg_val = avg
            elif agg_func.lower() == "min":
                agg_val = data["min_val"]
            elif agg_func.lower() == "max":
                agg_val = data["max_val"]
            elif agg_func.lower() == "count":
                agg_val = float(data["count"])
            else:
                agg_val = total

            results.append({
                "group": g_key,
                "aggregated_value": round(agg_val, 2),
                "total_sum": round(total, 2),
                "average": round(avg, 2),
                "count": data["count"],
                "max_cell": data["max_cell"],
                "max_value": round(data["max_val"], 2),
                "min_cell": data["min_cell"],
                "min_value": round(data["min_val"], 2),
                "sample_citation_cells": data["cells"][:3] + data["cells"][-3:] if len(data["cells"]) > 6 else data["cells"]
            })

        results.sort(key=lambda x: x["aggregated_value"], reverse=not ascending)
        if top_n > 0:
            results = results[:top_n]

        return {
            "sheet": ws.title,
            "group_by": group_by_col,
            "metric": metric_col,
            "agg_func": agg_func,
            "groups_count": len(results),
            "ranking": results
        }

    def add_calculated_column(self, header: str, formula_template: str, number_format: Optional[str] = FinancialStyleGuide.CURRENCY_FORMAT, sheet: Optional[str] = None, autofit: bool = True) -> Dict[str, Any]:
        from openpyxl.worksheet.table import TableColumn
        ws = self._get_sheet(sheet)
        new_col = ws.max_column + 1
        max_row = ws.max_row

        # Set header
        ws.cell(row=1, column=new_col, value=header)

        # Parse formula_template to adapt relative row numbers and sanitize compatibility
        sanitized_template = sanitize_formula_compatibility(formula_template)
        # If formula has row 2 (e.g. '=L2/J2') or {r}/{row}, adapt while strictly protecting absolute references ($)
        for r in range(2, max_row + 1):
            cell = ws.cell(row=r, column=new_col)
            row_formula = propagate_formula_row(sanitized_template, r, base_row=2)
            cell.value = row_formula
            if number_format:
                cell.number_format = number_format

        # Auto-resize openpyxl table if present
        table_resized = None
        for tbl in ws.tables.values():
            min_c, min_r, max_c, max_r = range_boundaries(tbl.ref)
            if new_col == max_c + 1:
                new_ref = f"{get_column_letter(min_c)}{min_r}:{get_column_letter(new_col)}{max_r}"
                tbl.ref = new_ref
                new_col_id = len(tbl.tableColumns) + 1
                tbl.tableColumns.append(TableColumn(id=new_col_id, name=header))
                table_resized = tbl.name
                break

        if autofit:
            col_letter = get_column_letter(new_col)
            max_len = max(len(str(ws.cell(row=r, column=new_col).value or '')) for r in range(1, max_row + 1))
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        return {
            "status": "success",
            "sheet": ws.title,
            "header": header,
            "column_index": new_col,
            "formula_applied": sanitized_template,
            "rows_affected": max_row - 1,
            "table_resized": table_resized
        }

    def apply_style_preset(self, range_str: str, preset: str, sheet: Optional[str] = None) -> Dict[str, Any]:
        """Aplica un preset de estilo universal a un rango en Headless OpenPyXL."""
        ws = self._get_sheet(sheet)
        if ":" in range_str:
            cells = [cell for row in ws[range_str] for cell in row]
        else:
            cells = [ws[range_str]]

        apply_style_preset_openpyxl(cells, preset)

        return {
            "status": "success",
            "range": range_str,
            "preset": preset,
            "sheet": ws.title,
            "cells_affected": len(cells)
        }

    def create_summary_table(
        self,
        title: str,
        headers: List[str],
        rows_spec: List[Dict[str, Any]],
        dest_sheet: str,
        start_cell: str = "B2",
        parameters: Optional[List[Dict[str, Any]]] = None
    ) -> Dict[str, Any]:
        """
        Crea una tabla resumen / comparativa profesional en Headless OpenPyXL con parámetros vinculados,
        encabezados corporativos, filas calculadas/estáticas, totales formateados y cuadrícula activa.
        """
        if dest_sheet in self.wb.sheetnames:
            ws = self.wb[dest_sheet]
        else:
            ws = self.wb.create_sheet(title=dest_sheet)

        # Activar visualización de cuadrícula
        if ws.views and ws.views.sheetView:
            ws.views.sheetView[0].showGridLines = True

        start_c_str, start_r_idx = coordinate_from_string(start_cell)
        start_col_idx = column_index_from_string(start_c_str)
        curr_r = start_r_idx

        # 1. Bloque de Parámetros de Entrada (si se especifican)
        param_cells: Dict[str, str] = {}
        if parameters:
            param_title = ws.cell(row=curr_r, column=start_col_idx, value="Supuestos y Parámetros")
            param_title.font = Font(name="Segoe UI", size=11, bold=True, color="0E2E63")
            curr_r += 1

            for param in parameters:
                param_name = param.get("name") or param.get("label", "Parámetro")
                param_val = param.get("value")
                param_format = param.get("format")
                param_note = param.get("note")

                if param.get("cell"):
                    val_coord = param["cell"]
                    v_cell = ws[val_coord]
                    v_cell.value = param_val
                else:
                    lbl_cell = ws.cell(row=curr_r, column=start_col_idx, value=param_name)
                    lbl_cell.font = Font(name="Segoe UI", size=10, bold=True)
                    v_cell = ws.cell(row=curr_r, column=start_col_idx + 1, value=param_val)
                    val_coord = v_cell.coordinate
                    curr_r += 1

                self.apply_style_preset(val_coord, "input_cell", sheet=dest_sheet)
                if param_format:
                    if param_format.lower() == "percent":
                        self.apply_style_preset(val_coord, "percent", sheet=dest_sheet)
                    elif param_format.lower() == "currency":
                        self.apply_style_preset(val_coord, "currency", sheet=dest_sheet)
                    else:
                        ws[val_coord].number_format = param_format

                if param_note:
                    ws[val_coord].comment = Comment(param_note, "Antigravity")

                param_cells[param_name] = val_coord

            curr_r += 1

        # 2. Título general de la tabla
        tbl_title_row = curr_r
        tbl_title_cell = ws.cell(row=tbl_title_row, column=start_col_idx, value=title)
        tbl_title_cell.font = Font(name="Segoe UI", size=13, bold=True, color="0E2E63")
        curr_r += 1

        # 3. Fila de encabezados
        header_row = curr_r
        end_col_idx = start_col_idx + len(headers) - 1
        for idx, h in enumerate(headers):
            ws.cell(row=header_row, column=start_col_idx + idx, value=h)

        hdr_range = f"{get_column_letter(start_col_idx)}{header_row}:{get_column_letter(end_col_idx)}{header_row}"
        self.apply_style_preset(hdr_range, "header", sheet=dest_sheet)
        curr_r += 1

        # 4. Filas de datos y totales
        for row_spec in rows_spec:
            is_total = row_spec.get("is_total", False) or row_spec.get("total", False)
            row_range = f"{get_column_letter(start_col_idx)}{curr_r}:{get_column_letter(end_col_idx)}{curr_r}"

            label = row_spec.get("label") or row_spec.get("name") or row_spec.get("title")
            row_vals = []
            if "values" in row_spec:
                vals = row_spec["values"]
                if len(vals) == len(headers):
                    row_vals = vals
                elif len(vals) == len(headers) - 1:
                    row_vals = [label] + vals
                else:
                    row_vals = vals
            elif "formulas" in row_spec:
                forms = row_spec["formulas"]
                if len(forms) == len(headers):
                    row_vals = forms
                elif len(forms) == len(headers) - 1:
                    row_vals = [label] + forms
                else:
                    row_vals = forms
            else:
                row_vals = [row_spec.get(h, "") for h in headers]
                if label and not row_vals[0]:
                    row_vals[0] = label

            formats = row_spec.get("formats")
            row_fmt = row_spec.get("format") or row_spec.get("number_format")
            presets = row_spec.get("presets") or row_spec.get("styles")

            for c_idx, val in enumerate(row_vals):
                if c_idx >= len(headers):
                    break
                cell = ws.cell(row=curr_r, column=start_col_idx + c_idx)
                if isinstance(val, str) and val.startswith("="):
                    cell.value = sanitize_formula_compatibility(val)
                else:
                    cell.value = val

                if not is_total:
                    cell.font = Font(name="Segoe UI", size=10)
                    cell.border = Border(
                        left=Side(style='thin', color='D9D9D9'),
                        right=Side(style='thin', color='D9D9D9'),
                        top=Side(style='thin', color='D9D9D9'),
                        bottom=Side(style='thin', color='D9D9D9')
                    )

                col_fmt = formats[c_idx] if (formats and c_idx < len(formats)) else row_fmt
                if col_fmt:
                    if col_fmt.lower() == "currency":
                        cell.number_format = FinancialStyleGuide.CURRENCY_FORMAT
                    elif col_fmt.lower() == "percent":
                        cell.number_format = FinancialStyleGuide.PERCENT_FORMAT
                    else:
                        cell.number_format = col_fmt

                if presets and c_idx < len(presets) and presets[c_idx]:
                    self.apply_style_preset(cell.coordinate, presets[c_idx], sheet=dest_sheet)

            if is_total:
                self.apply_style_preset(row_range, "total_row", sheet=dest_sheet)

            curr_r += 1

        # 5. AutoFit en columnas involucradas
        for c_idx in range(start_col_idx, end_col_idx + 1):
            col_letter = get_column_letter(c_idx)
            max_len = max(len(str(ws.cell(row=r, column=c_idx).value or '')) for r in range(start_r_idx, curr_r))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        return {
            "status": "success",
            "sheet": dest_sheet,
            "title": title,
            "table_range": f"{get_column_letter(start_col_idx)}{header_row}:{get_column_letter(end_col_idx)}{curr_r - 1}",
            "rows_count": len(rows_spec),
            "headers": headers,
            "parameters": param_cells,
            "start_cell": start_cell
        }

    def create_pivot_table(
        self,
        source: Union[str, Dict[str, Any]],
        rows: List[str],
        cols: Optional[List[str]] = None,
        value_field: Optional[str] = None,
        value_func: str = "sum",
        value_format: Optional[str] = FinancialStyleGuide.CURRENCY_FORMAT,
        dest_sheet: Optional[str] = None,
        dest_cell: str = "A3",
        table_name: Optional[str] = None,
        source_sheet: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        Crea una hoja de análisis con tabla de doble entrada cruzada (cross-tabulation resumida)
        con totales generales formateados en Headless OpenPyXL.
        """
        src_ws = self._get_sheet(source_sheet) if source_sheet else self.wb.active

        # 1. Resolver rango origen y encabezados
        if isinstance(source, dict) and "range" in source:
            source_range_str = source["range"]
        else:
            source_range_str = str(source)

        min_col, min_row, max_col, max_row = 1, 1, src_ws.max_column, src_ws.max_row
        if ":" in source_range_str:
            min_col, min_row, max_col, max_row = range_boundaries(source_range_str)
        elif source_range_str in src_ws.tables:
            min_col, min_row, max_col, max_row = range_boundaries(src_ws.tables[source_range_str].ref)

        headers = {}
        for c in range(min_col, max_col + 1):
            val = src_ws.cell(row=min_row, column=c).value
            if val is not None:
                headers[str(val).strip().lower()] = c

        row_key_name = rows[0] if rows else ""
        col_key_name = cols[0] if cols else None
        metric_name = value_field or ""

        row_c_idx = headers.get(row_key_name.strip().lower())
        col_c_idx = headers.get(col_key_name.strip().lower()) if col_key_name else None
        val_c_idx = headers.get(metric_name.strip().lower())

        if not row_c_idx:
            raise ValueError(f"Campo de fila '{row_key_name}' no encontrado en encabezados.")
        if not val_c_idx:
            raise ValueError(f"Campo de valor '{metric_name}' no encontrado en encabezados.")

        # 2. Agregar datos cruzados en memoria
        # data_matrix[row_val][col_val] = list of floats
        data_matrix: Dict[str, Dict[str, List[float]]] = {}
        row_values_ordered = []
        col_values_ordered = []

        for r in range(min_row + 1, max_row + 1):
            r_val = src_ws.cell(row=r, column=row_c_idx).value
            v_val = src_ws.cell(row=r, column=val_c_idx).value
            c_val = src_ws.cell(row=r, column=col_c_idx).value if col_c_idx else "Total"

            if r_val is None:
                continue
            try:
                num = float(v_val)
            except (ValueError, TypeError):
                continue

            r_str = str(r_val).strip()
            c_str = str(c_val).strip()

            if r_str not in data_matrix:
                data_matrix[r_str] = {}
                row_values_ordered.append(r_str)
            if c_str not in data_matrix[r_str]:
                data_matrix[r_str][c_str] = []
            if c_str not in col_values_ordered:
                col_values_ordered.append(c_str)

            data_matrix[r_str][c_str].append(num)

        # 3. Preparar hoja de destino
        target_sheet_name = dest_sheet or (src_ws.title + "_Pivot")
        if target_sheet_name in self.wb.sheetnames:
            dest_ws = self.wb[target_sheet_name]
        else:
            dest_ws = self.wb.create_sheet(title=target_sheet_name)

        # Parse dest_cell
        start_c_idx, start_r_idx = coordinate_from_string(dest_cell)
        start_col_num = column_index_from_string(start_c_idx)

        # 4. Escribir tabla cruzada
        # Estilos corporativos
        header_fill = PatternFill(start_color="0E2E63", end_color="0E2E63", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        total_font = Font(name="Segoe UI", size=10, bold=True)
        data_font = Font(name="Segoe UI", size=10)
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        total_border = Border(
            top=Side(style='thin', color='000000'),
            bottom=Side(style='double', color='000000')
        )

        def calc_agg(nums: List[float]) -> float:
            if not nums:
                return 0.0
            if value_func.lower() == "sum":
                return sum(nums)
            elif value_func.lower() in ("avg", "average", "mean"):
                return sum(nums) / len(nums)
            elif value_func.lower() == "count":
                return float(len(nums))
            elif value_func.lower() == "min":
                return min(nums)
            elif value_func.lower() == "max":
                return max(nums)
            return sum(nums)

        # Escribir encabezados de columnas
        # Fila de encabezado: [Row_Field_Name, Col1, Col2, ..., Total General]
        header_r = start_r_idx
        # Celda de fila
        c_cell = dest_ws.cell(row=header_r, column=start_col_num, value=row_key_name)
        c_cell.font = header_font
        c_cell.fill = header_fill

        num_data_cols = len(col_values_ordered)
        for idx, col_val in enumerate(col_values_ordered):
            cell = dest_ws.cell(row=header_r, column=start_col_num + 1 + idx, value=col_val)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="right")

        # Columna Total General
        tot_col_idx = start_col_num + 1 + num_data_cols
        tot_header_cell = dest_ws.cell(row=header_r, column=tot_col_idx, value="Total General")
        tot_header_cell.font = header_font
        tot_header_cell.fill = header_fill
        tot_header_cell.alignment = Alignment(horizontal="right")

        # Escribir filas de datos
        curr_r = header_r + 1
        col_accumulators: Dict[str, List[float]] = {c: [] for c in col_values_ordered}
        all_nums: List[float] = []

        for r_val in row_values_ordered:
            row_label_cell = dest_ws.cell(row=curr_r, column=start_col_num, value=r_val)
            row_label_cell.font = data_font
            row_label_cell.border = thin_border

            row_nums = []
            for idx, col_val in enumerate(col_values_ordered):
                nums = data_matrix[r_val].get(col_val, [])
                agg_res = calc_agg(nums)
                cell = dest_ws.cell(row=curr_r, column=start_col_num + 1 + idx, value=agg_res)
                cell.font = data_font
                cell.border = thin_border
                if value_format:
                    cell.number_format = value_format
                if nums:
                    col_accumulators[col_val].extend(nums)
                    row_nums.extend(nums)

            # Total de la fila
            row_tot = calc_agg(row_nums)
            tot_cell = dest_ws.cell(row=curr_r, column=tot_col_idx, value=row_tot)
            tot_cell.font = total_font
            tot_cell.border = thin_border
            if value_format:
                tot_cell.number_format = value_format
            all_nums.extend(row_nums)
            curr_r += 1

        # Fila de Total General al pie
        footer_r = curr_r
        foot_label_cell = dest_ws.cell(row=footer_r, column=start_col_num, value="Total General")
        foot_label_cell.font = total_font
        foot_label_cell.border = total_border

        for idx, col_val in enumerate(col_values_ordered):
            col_tot = calc_agg(col_accumulators[col_val])
            cell = dest_ws.cell(row=footer_r, column=start_col_num + 1 + idx, value=col_tot)
            cell.font = total_font
            cell.border = total_border
            if value_format:
                cell.number_format = value_format

        # Total absoluto
        grand_total = calc_agg(all_nums)
        grand_tot_cell = dest_ws.cell(row=footer_r, column=tot_col_idx, value=grand_total)
        grand_tot_cell.font = total_font
        grand_tot_cell.border = total_border
        if value_format:
            grand_tot_cell.number_format = value_format

        # AutoFit en columnas de la hoja destino
        total_cols_span = tot_col_idx - start_col_num + 1
        for col_idx in range(start_col_num, tot_col_idx + 1):
            col_letter = get_column_letter(col_idx)
            max_len = max(len(str(dest_ws.cell(row=r, column=col_idx).value or '')) for r in range(header_r, footer_r + 1))
            dest_ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        end_cell = f"{get_column_letter(tot_col_idx)}{footer_r}"
        pt_range = f"{dest_cell}:{end_cell}"

        return {
            "status": "success",
            "pivot_table_name": table_name or "Pivot_Summary",
            "dest_sheet": dest_ws.title,
            "dest_range": pt_range,
            "rows": rows,
            "cols": cols or [],
            "value_field": metric_name,
            "value_func": value_func,
            "value_format": value_format,
            "grand_total": round(grand_total, 2)
        }

    def create_chart(
        self,
        source: Union[str, Dict[str, Any]],
        chart_type: str = "column_clustered",
        title: Optional[str] = None,
        dest_sheet: Optional[str] = None,
        placement: str = "auto",
        anchor: Optional[str] = None,
        width: int = 480,
        height: int = 300,
        source_sheet: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        Inserta un gráfico nativo openpyxl.chart en la posición calculada anti-solapamiento.
        """
        from openpyxl.chart import BarChart, LineChart, PieChart, AreaChart, Reference

        target_ws = self._get_sheet(dest_sheet) if dest_sheet else self.wb.active

        # Resolver rango fuente
        if isinstance(source, dict):
            if "dest_sheet" in source and not dest_sheet:
                target_ws = self._get_sheet(source["dest_sheet"])
            source_range_str = source.get("dest_range", "A3:B10")
        else:
            source_range_str = str(source)

        min_col, min_row, max_col, max_row = range_boundaries(source_range_str)

        # 1. Instanciar clase de gráfico correspondiente
        ct = chart_type.lower()
        if ct in ("column_clustered", "column"):
            chart = BarChart()
            chart.type = "col"
            chart.grouping = "clustered"
        elif ct in ("bar_clustered", "bar"):
            chart = BarChart()
            chart.type = "bar"
            chart.grouping = "clustered"
        elif ct == "line":
            chart = LineChart()
        elif ct == "pie":
            chart = PieChart()
        elif ct == "area":
            chart = AreaChart()
        else:
            chart = BarChart()
            chart.type = "col"

        # 2. Configurar título y dimensiones
        if title:
            chart.title = title
        chart.width = width / 30.0 # openpyxl usa cm aprox (480px ~ 16cm)
        chart.height = height / 30.0 # (300px ~ 10cm)

        # 3. Referencias de datos y categorías
        # Suponiendo que la primera columna son categorías y la última fila puede ser total general
        # Si la última fila es 'Total General', excluirla de la serie para no desvirtuar el gráfico
        last_row_label = str(target_ws.cell(row=max_row, column=min_col).value or "").strip().lower()
        data_end_row = max_row - 1 if last_row_label == "total general" else max_row

        # Si hay más de 2 columnas y la última columna es 'Total General', excluirla si hay desglose por columnas
        last_col_header = str(target_ws.cell(row=min_row, column=max_col).value or "").strip().lower()
        data_end_col = max_col - 1 if (last_col_header == "total general" and max_col - min_col >= 2) else max_col

        data_ref = Reference(target_ws, min_col=min_col + 1, min_row=min_row, max_col=data_end_col, max_row=data_end_row)
        cats_ref = Reference(target_ws, min_col=min_col, min_row=min_row + 1, max_row=data_end_row)

        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)

        # 4. Calcular posición geométrica anti-solapamiento
        num_cols = max_col - min_col + 1
        num_rows = max_row - min_row + 1
        pos = self._calculate_placement_coordinates(
            start_col=min_col,
            start_row=min_row,
            num_cols=num_cols,
            num_rows=num_rows,
            placement=placement,
            anchor=anchor
        )

        anchor_cell = pos["cell"]
        target_ws.add_chart(chart, anchor_cell)

        return {
            "status": "success",
            "sheet": target_ws.title,
            "chart_type": chart_type,
            "title": title,
            "placement": pos["placement"],
            "anchor_cell": anchor_cell,
            "data_range": f"{get_column_letter(min_col + 1)}{min_row}:{get_column_letter(data_end_col)}{data_end_row}"
        }

    def save(self, file_path: Optional[str] = None) -> str:
        out_path = file_path or self.file_path
        if out_path:
            self.wb.save(out_path)
            return os.path.abspath(out_path)
        raise ValueError("No se especificó ruta para guardar el archivo.")

    def remove_duplicates(self, range_or_table: str, key_columns: Optional[List[str]] = None, sheet: Optional[str] = None) -> Dict[str, Any]:
        """
        Elimina filas duplicadas dentro del rango o tabla estructurada on-demand.
        En OpenPyXL, elimina filas y redimensiona tbl.ref en la tabla estructurada si existe.
        """
        ws = self._get_sheet(sheet)
        tbl = None
        if range_or_table and range_or_table in ws.tables:
            tbl = ws.tables[range_or_table]
            min_col, min_row, max_col, max_row = range_boundaries(tbl.ref)
        elif range_or_table:
            min_col, min_row, max_col, max_row = range_boundaries(range_or_table)
            for t in ws.tables.values():
                if t.ref == range_or_table:
                    tbl = t
                    break
        elif ws.tables:
            tbl = next(iter(ws.tables.values()))
            min_col, min_row, max_col, max_row = range_boundaries(tbl.ref)
        else:
            min_col, min_row, max_col, max_row = 1, 1, ws.max_column, ws.max_row

        num_cols = max_col - min_col + 1
        data_rows_count = max_row - min_row
        if data_rows_count <= 0:
            return {
                "status": "success",
                "range_or_table": range_or_table or (tbl.name if tbl else f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"),
                "sheet": ws.title,
                "original_rows": 0,
                "duplicates_removed": 0,
                "remaining_rows": 0,
                "key_columns": key_columns or "all"
            }

        headers = [str(ws.cell(row=min_row, column=c).value or "").strip() for c in range(min_col, max_col + 1)]
        headers_lower = [h.lower() for h in headers]

        compare_cols = []
        if key_columns:
            for k in key_columns:
                k_clean = str(k).strip()
                k_lower = k_clean.lower()
                if k_lower in headers_lower:
                    idx = headers_lower.index(k_lower) + 1
                    compare_cols.append(idx)
                elif k_clean.isalpha() and len(k_clean) <= 3:
                    c_idx = column_index_from_string(k_clean)
                    rel_idx = c_idx - min_col + 1
                    if 1 <= rel_idx <= num_cols:
                        compare_cols.append(rel_idx)
            if not compare_cols:
                compare_cols = list(range(1, num_cols + 1))
        else:
            compare_cols = list(range(1, num_cols + 1))

        seen = set()
        dup_row_indices = []

        for r in range(min_row + 1, max_row + 1):
            row_vals = [ws.cell(row=r, column=min_col + c - 1).value for c in range(1, num_cols + 1)]
            key = tuple(row_vals[ci - 1] for ci in compare_cols)
            if key in seen:
                dup_row_indices.append(r)
            else:
                seen.add(key)

        duplicates_count = len(dup_row_indices)
        if duplicates_count > 0:
            for r in sorted(dup_row_indices, reverse=True):
                ws.delete_rows(r, 1)

            if tbl:
                new_max_row = max_row - duplicates_count
                tbl.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{new_max_row}"

        remaining_data_rows = data_rows_count - duplicates_count
        return {
            "status": "success",
            "range_or_table": range_or_table or (tbl.name if tbl else f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"),
            "sheet": ws.title,
            "original_rows": data_rows_count,
            "duplicates_removed": duplicates_count,
            "remaining_rows": remaining_data_rows,
            "key_columns": key_columns or "all"
        }

    def detect_and_flag_outliers(
        self,
        column: str,
        method: str = "iqr",
        range_or_table: Optional[str] = None,
        sheet: Optional[str] = None,
        flag_style: str = "warning",
        add_comment: bool = True
    ) -> Dict[str, Any]:
        """
        Detecta y señala visualmente valores atípicos (outliers) mediante IQR o Z-score on-demand.
        Aplica formato no destructivo (#FCE4D6 y texto rojo bold) y añade comentarios a las celdas.
        """
        ws = self._get_sheet(sheet)
        tbl = None
        if range_or_table and range_or_table in ws.tables:
            tbl = ws.tables[range_or_table]
            min_col, min_row, max_col, max_row = range_boundaries(tbl.ref)
        elif range_or_table:
            min_col, min_row, max_col, max_row = range_boundaries(range_or_table)
            for t in ws.tables.values():
                if t.ref == range_or_table:
                    tbl = t
                    break
        elif ws.tables:
            tbl = next(iter(ws.tables.values()))
            min_col, min_row, max_col, max_row = range_boundaries(tbl.ref)
        else:
            min_col, min_row, max_col, max_row = 1, 1, ws.max_column, ws.max_row

        num_cols = max_col - min_col + 1
        headers = [str(ws.cell(row=min_row, column=c).value or "").strip() for c in range(min_col, max_col + 1)]
        headers_lower = [h.lower() for h in headers]
        col_clean = str(column).strip()
        col_lower = col_clean.lower()

        target_col_rel = None
        if col_lower in headers_lower:
            target_col_rel = headers_lower.index(col_lower) + 1
        elif col_clean.isalpha() and len(col_clean) <= 3:
            c_idx = column_index_from_string(col_clean)
            rel_idx = c_idx - min_col + 1
            if 1 <= rel_idx <= num_cols:
                target_col_rel = rel_idx

        if not target_col_rel:
            raise ValueError(f"Columna '{column}' no encontrada en el rango/tabla. Encabezados disponibles: {headers}")

        abs_col_idx = min_col + target_col_rel - 1

        val_list = []
        row_val_map = {}
        for r in range(min_row + 1, max_row + 1):
            v = ws.cell(row=r, column=abs_col_idx).value
            if v is not None and str(v).strip() != "":
                try:
                    fv = float(str(v).replace("$", "").replace(" ", "").replace(",", ""))
                    val_list.append(fv)
                    row_val_map[r] = fv
                except Exception:
                    pass

        if len(val_list) < 4:
            return {
                "status": "success",
                "sheet": ws.title,
                "column": column,
                "method": method,
                "total_values_evaluated": len(val_list),
                "outliers_count": 0,
                "thresholds": {},
                "flagged_cells": [],
                "outliers": []
            }

        method_clean = method.lower()
        if method_clean == "zscore":
            mean = statistics.mean(val_list)
            std = statistics.stdev(val_list) if len(val_list) > 1 else 0
            z_thresh = 2.5 if len(val_list) < 30 else 3.0
            lower_bound = mean - z_thresh * std
            upper_bound = mean + z_thresh * std
            thresholds = {"lower_bound": lower_bound, "upper_bound": upper_bound, "mean": mean, "std": std, "z_threshold": z_thresh}
        else:
            sorted_vals = sorted(val_list)
            q1, _, q3 = statistics.quantiles(sorted_vals, n=4, method="inclusive")
            iqr = q3 - q1
            lower_bound = q1 - 1.5 * iqr
            upper_bound = q3 + 1.5 * iqr
            thresholds = {"lower_bound": lower_bound, "upper_bound": upper_bound, "q1": q1, "q3": q3, "iqr": iqr}

        outliers = []
        col_letter = get_column_letter(abs_col_idx)

        warning_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        warning_font = Font(color="C00000", bold=True)

        for r, fv in row_val_map.items():
            if fv < lower_bound or fv > upper_bound:
                cell = ws.cell(row=r, column=abs_col_idx)
                cell_coord = f"{col_letter}{r}"
                cell.fill = warning_fill
                cell.font = warning_font

                if add_comment:
                    try:
                        from openpyxl.comments import Comment
                        cell.comment = Comment(
                            f"Outlier detectado ({method_clean.upper()}): valor={fv}, límites=[{lower_bound:.2f}, {upper_bound:.2f}]",
                            "Antigravity"
                        )
                    except Exception:
                        pass

                outliers.append({
                    "cell": cell_coord,
                    "row": r,
                    "value": fv,
                    "reason": f"Valor {fv} fuera de límites [{lower_bound:.2f}, {upper_bound:.2f}]"
                })

        return {
            "status": "success",
            "sheet": ws.title,
            "column": headers[target_col_rel - 1],
            "method": method_clean,
            "total_values_evaluated": len(val_list),
            "outliers_count": len(outliers),
            "thresholds": thresholds,
            "flagged_cells": [o["cell"] for o in outliers],
            "outliers": outliers
        }

    def clean_table_dataset(self, target_range_or_table: str, rules: Dict[str, Any], sheet: Optional[str] = None) -> Dict[str, Any]:
        """
        Limpia y estandariza en bloque los datos de una tabla o rango on-demand.
        Procesa en una sola pasada en memoria:
          - standardize_text: trim, casing y mapeo de sinónimos/reemplazos.
          - cast_types: conversión segura de tipos (float, int, str).
          - coerce_dates: conversión a seriales de Excel o datetimes con formato uniforme.
        """
        ws = self._get_sheet(sheet)
        tbl = None
        if target_range_or_table and target_range_or_table in ws.tables:
            tbl = ws.tables[target_range_or_table]
            min_col, min_row, max_col, max_row = range_boundaries(tbl.ref)
        elif target_range_or_table:
            min_col, min_row, max_col, max_row = range_boundaries(target_range_or_table)
            for t in ws.tables.values():
                if t.ref == target_range_or_table:
                    tbl = t
                    break
        elif ws.tables:
            tbl = next(iter(ws.tables.values()))
            min_col, min_row, max_col, max_row = range_boundaries(tbl.ref)
        else:
            min_col, min_row, max_col, max_row = 1, 1, ws.max_column, ws.max_row

        num_cols = max_col - min_col + 1
        data_rows_count = max_row - min_row

        headers = [str(ws.cell(row=min_row, column=c).value or "").strip() for c in range(min_col, max_col + 1)]

        std_text_rules = rules.get("standardize_text", {})
        cast_type_rules = rules.get("cast_types", {})
        coerce_date_rules = rules.get("coerce_dates", {})

        col_text_cfg = {}
        col_cast_cfg = {}
        col_date_cfg = {}

        for c_idx, h in enumerate(headers, start=1):
            h_lower = h.lower()
            for rule_col, cfg in std_text_rules.items():
                if rule_col.strip().lower() == h_lower:
                    col_text_cfg[c_idx] = cfg
                    break

            for rule_col, target_t in cast_type_rules.items():
                if rule_col.strip().lower() == h_lower:
                    col_cast_cfg[c_idx] = str(target_t).lower()
                    break

            if isinstance(coerce_date_rules, list):
                if any(dc.strip().lower() == h_lower for dc in coerce_date_rules):
                    col_date_cfg[c_idx] = {"to_serial": True, "format": "yyyy-mm-dd"}
            elif isinstance(coerce_date_rules, dict):
                for dc, cfg in coerce_date_rules.items():
                    if dc.strip().lower() == h_lower:
                        col_date_cfg[c_idx] = cfg if isinstance(cfg, dict) else {"to_serial": True, "format": "yyyy-mm-dd"}
                        break

        cells_modified = 0

        if rules.get("trim_headers", True):
            for c_idx in range(1, num_cols + 1):
                cell = ws.cell(row=min_row, column=min_col + c_idx - 1)
                if isinstance(cell.value, str) and cell.value != cell.value.strip():
                    cell.value = cell.value.strip()
                    cells_modified += 1

        for r in range(min_row + 1, max_row + 1):
            for c_idx in range(1, num_cols + 1):
                cell = ws.cell(row=r, column=min_col + c_idx - 1)
                val = cell.value
                orig_val = val

                # 1. Date coercion
                if c_idx in col_date_cfg:
                    serial = parse_excel_date_serial(val)
                    if serial is not None:
                        val = serial
                        d_cfg = col_date_cfg[c_idx]
                        fmt = d_cfg.get("format", "yyyy-mm-dd") if isinstance(d_cfg, dict) else "yyyy-mm-dd"
                        cell.number_format = fmt

                # 2. Text standardization
                if c_idx in col_text_cfg and val is not None:
                    t_cfg = col_text_cfg[c_idx]
                    s = str(val)
                    if t_cfg.get("trim", True):
                        s = s.strip()
                    reps = t_cfg.get("replacements") or t_cfg.get("mapping")
                    if reps:
                        if s in reps:
                            s = reps[s]
                        elif s.lower() in {k.lower(): v for k, v in reps.items()}:
                            for k, v in reps.items():
                                if s.lower() == k.lower():
                                    s = v
                                    break
                    casing = t_cfg.get("casing") or t_cfg.get("case")
                    if casing == "upper":
                        s = s.upper()
                    elif casing == "lower":
                        s = s.lower()
                    elif casing == "title":
                        s = s.title()
                    elif casing == "capitalize":
                        s = s.capitalize()
                    val = s

                # 3. Type casting
                if c_idx in col_cast_cfg and val is not None and str(val).strip() != "":
                    tgt = col_cast_cfg[c_idx]
                    try:
                        clean_str = str(val).replace("$", "").replace(" ", "").replace(",", "")
                        if tgt in ("float", "double", "decimal"):
                            val = float(clean_str)
                        elif tgt in ("int", "integer"):
                            val = int(round(float(clean_str)))
                        elif tgt in ("str", "string", "text"):
                            val = str(val)
                    except Exception:
                        pass

                if val != orig_val:
                    cell.value = val
                    cells_modified += 1

        return {
            "status": "success",
            "sheet": ws.title,
            "target_range_or_table": target_range_or_table or (tbl.name if tbl else f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"),
            "rows_processed": data_rows_count,
            "cells_modified": cells_modified,
            "rules_applied": list(rules.keys())
        }

    def curate_pipeline(
        self,
        target_range_or_table: str,
        pipeline_spec: Dict[str, Any],
        sheet: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        Pipeline unificado de curaduría y vectorización en memoria para OpenPyXL headless.
        Ejecuta deduplicación, estandarización de texto, casteo de tipos, coerción de fechas
        y detección/marcado de outliers en una sola pasada en memoria.
        """
        ws = self._get_sheet(sheet)
        tbl = None
        if target_range_or_table and target_range_or_table in ws.tables:
            tbl = ws.tables[target_range_or_table]
            min_col, min_row, max_col, max_row = range_boundaries(tbl.ref)
        elif target_range_or_table:
            min_col, min_row, max_col, max_row = range_boundaries(target_range_or_table)
            for t in ws.tables.values():
                if t.ref == target_range_or_table:
                    tbl = t
                    break
        elif ws.tables:
            tbl = next(iter(ws.tables.values()))
            min_col, min_row, max_col, max_row = range_boundaries(tbl.ref)
        else:
            min_col, min_row, max_col, max_row = 1, 1, ws.max_column, ws.max_row

        num_cols = max_col - min_col + 1
        data_rows_count = max_row - min_row
        if data_rows_count <= 0:
            return {
                "status": "success",
                "sheet": ws.title,
                "target_range_or_table": target_range_or_table or (tbl.name if tbl else f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"),
                "original_rows": 0,
                "remaining_rows": 0,
                "duplicates_removed": 0,
                "cells_cleaned": 0,
                "rules_applied": [],
                "outliers_count": 0,
                "outliers_flagged": [],
                "outliers_summary": []
            }

        # 1. Lectura a memoria
        headers = [str(ws.cell(row=min_row, column=c).value or "").strip() for c in range(min_col, max_col + 1)]
        data_rows = []
        for r in range(min_row + 1, max_row + 1):
            row_vals = [ws.cell(row=r, column=min_col + c - 1).value for c in range(1, num_cols + 1)]
            data_rows.append(row_vals)
        orig_data_rows_count = len(data_rows)

        # 2. Deduplicación en memoria
        dedup_spec = pipeline_spec.get("dedup")
        duplicates_count = 0
        if dedup_spec:
            if isinstance(dedup_spec, dict):
                key_cols = dedup_spec.get("key_columns")
            elif isinstance(dedup_spec, (list, tuple)):
                key_cols = list(dedup_spec)
            else:
                key_cols = None

            headers_lower = [h.lower() for h in headers]
            compare_cols = []
            if key_cols:
                for k in key_cols:
                    k_clean = str(k).strip()
                    k_lower = k_clean.lower()
                    if k_lower in headers_lower:
                        compare_cols.append(headers_lower.index(k_lower) + 1)
                    elif k_clean.isalpha() and len(k_clean) <= 3:
                        c_idx = column_index_from_string(k_clean)
                        rel_idx = c_idx - min_col + 1
                        if 1 <= rel_idx <= num_cols:
                            compare_cols.append(rel_idx)
                if not compare_cols:
                    compare_cols = list(range(1, num_cols + 1))
            else:
                compare_cols = list(range(1, num_cols + 1))

            seen = set()
            unique_rows = []
            for r in data_rows:
                key = tuple(
                    str(r[ci - 1]).strip() if isinstance(r[ci - 1], str) else r[ci - 1]
                    if ci - 1 < len(r) else None
                    for ci in compare_cols
                )
                if key not in seen:
                    seen.add(key)
                    unique_rows.append(r)
            duplicates_count = len(data_rows) - len(unique_rows)
            data_rows = unique_rows

        # 3. Limpieza y estandarización en memoria
        clean_spec = pipeline_spec.get("clean")
        cells_modified = 0
        col_date_cfg = {}
        if clean_spec:
            std_text_rules = clean_spec.get("standardize_text", {})
            cast_type_rules = clean_spec.get("cast_types", {})
            coerce_date_rules = clean_spec.get("coerce_dates", {})
            trim_headers = clean_spec.get("trim_headers", True)

            col_text_cfg = {}
            col_cast_cfg = {}

            for c_idx, h in enumerate(headers, start=1):
                h_lower = h.lower()
                for rule_col, cfg in std_text_rules.items():
                    if rule_col.strip().lower() == h_lower:
                        col_text_cfg[c_idx] = cfg
                        break
                for rule_col, tgt_t in cast_type_rules.items():
                    if rule_col.strip().lower() == h_lower:
                        col_cast_cfg[c_idx] = str(tgt_t).lower()
                        break
                if isinstance(coerce_date_rules, (list, tuple)):
                    if any(dc.strip().lower() == h_lower for dc in coerce_date_rules):
                        col_date_cfg[c_idx] = {"to_serial": True, "format": "yyyy-mm-dd"}
                elif isinstance(coerce_date_rules, dict):
                    for dc, cfg in coerce_date_rules.items():
                        if dc.strip().lower() == h_lower:
                            col_date_cfg[c_idx] = cfg if isinstance(cfg, dict) else {"to_serial": True, "format": "yyyy-mm-dd"}
                            break

            if trim_headers:
                new_headers = []
                for h in headers:
                    th = h.strip() if isinstance(h, str) else h
                    if th != h:
                        cells_modified += 1
                    new_headers.append(th)
                headers = new_headers

            cleaned_data_rows = []
            for r in data_rows:
                new_row = []
                for c_idx in range(1, num_cols + 1):
                    val = r[c_idx - 1] if c_idx - 1 < len(r) else None
                    orig_val = val

                    # 1. Coerción de fechas
                    if c_idx in col_date_cfg:
                        serial = parse_excel_date_serial(val)
                        if serial is not None:
                            val = serial

                    # 2. Estandarización de texto
                    if c_idx in col_text_cfg and val is not None:
                        t_cfg = col_text_cfg[c_idx]
                        s = str(val)
                        if t_cfg.get("trim", True):
                            s = s.strip()
                        reps = t_cfg.get("replacements") or t_cfg.get("mapping")
                        if reps:
                            if s in reps:
                                s = reps[s]
                            elif s.lower() in {k.lower(): v for k, v in reps.items()}:
                                for k, v in reps.items():
                                    if s.lower() == k.lower():
                                        s = v
                                        break
                        casing = t_cfg.get("casing") or t_cfg.get("case")
                        if casing == "upper":
                            s = s.upper()
                        elif casing == "lower":
                            s = s.lower()
                        elif casing == "title":
                            s = s.title()
                        elif casing == "capitalize":
                            s = s.capitalize()
                        val = s

                    # 3. Casteo seguro de tipos
                    if c_idx in col_cast_cfg and val is not None and str(val).strip() != "":
                        tgt = col_cast_cfg[c_idx]
                        try:
                            clean_str = str(val).replace("$", "").replace(" ", "").replace(",", "")
                            if tgt in ("float", "double", "decimal"):
                                val = float(clean_str)
                            elif tgt in ("int", "integer"):
                                val = int(round(float(clean_str)))
                            elif tgt in ("str", "string", "text"):
                                val = str(val)
                        except Exception:
                            pass

                    if val != orig_val:
                        cells_modified += 1
                    new_row.append(val)
                cleaned_data_rows.append(new_row)
            data_rows = cleaned_data_rows

        # 4. Cálculo estadístico de outliers en memoria
        outliers_spec = pipeline_spec.get("outliers")
        outliers_results = []
        all_flagged_cells = []
        outlier_cells_to_format = []

        if outliers_spec:
            if isinstance(outliers_spec, dict):
                if "column" in outliers_spec or "columns" in outliers_spec:
                    spec_list = [outliers_spec]
                else:
                    spec_list = [{"column": k, **(v if isinstance(v, dict) else {})} for k, v in outliers_spec.items()]
            elif isinstance(outliers_spec, (list, tuple)):
                spec_list = []
                for it in outliers_spec:
                    if isinstance(it, str):
                        spec_list.append({"column": it})
                    elif isinstance(it, dict):
                        spec_list.append(it)
            else:
                spec_list = []

            headers_lower = [h.lower() for h in headers]
            for s_item in spec_list:
                target_col_name = str(s_item.get("column") or "").strip()
                method = str(s_item.get("method") or "iqr").lower()
                add_comment = bool(s_item.get("add_comment", True))

                target_col_rel = None
                if target_col_name.lower() in headers_lower:
                    target_col_rel = headers_lower.index(target_col_name.lower()) + 1
                elif target_col_name.isalpha() and len(target_col_name) <= 3:
                    c_idx = column_index_from_string(target_col_name)
                    rel_idx = c_idx - min_col + 1
                    if 1 <= rel_idx <= num_cols:
                        target_col_rel = rel_idx

                if not target_col_rel:
                    continue

                abs_col_idx = min_col + target_col_rel - 1
                col_letter = get_column_letter(abs_col_idx)

                val_list = []
                row_val_map = {}
                for r_offset, r_data in enumerate(data_rows):
                    sheet_r = min_row + 1 + r_offset
                    v = r_data[target_col_rel - 1] if target_col_rel - 1 < len(r_data) else None
                    if v is not None and str(v).strip() != "":
                        try:
                            fv = float(str(v).replace("$", "").replace(" ", "").replace(",", ""))
                            val_list.append(fv)
                            row_val_map[sheet_r] = fv
                        except Exception:
                            pass

                if len(val_list) < 4:
                    outliers_results.append({
                        "column": headers[target_col_rel - 1],
                        "method": method,
                        "total_values_evaluated": len(val_list),
                        "outliers_count": 0,
                        "thresholds": {},
                        "flagged_cells": [],
                        "outliers": []
                    })
                    continue

                if method == "zscore":
                    mean = statistics.mean(val_list)
                    std = statistics.stdev(val_list) if len(val_list) > 1 else 0
                    z_thresh = 2.5 if len(val_list) < 30 else 3.0
                    lower_bound = mean - z_thresh * std
                    upper_bound = mean + z_thresh * std
                    thresholds = {"lower_bound": lower_bound, "upper_bound": upper_bound, "mean": mean, "std": std, "z_threshold": z_thresh}
                else:
                    sorted_vals = sorted(val_list)
                    q1, _, q3 = statistics.quantiles(sorted_vals, n=4, method="inclusive")
                    iqr = q3 - q1
                    lower_bound = q1 - 1.5 * iqr
                    upper_bound = q3 + 1.5 * iqr
                    thresholds = {"lower_bound": lower_bound, "upper_bound": upper_bound, "q1": q1, "q3": q3, "iqr": iqr}

                col_outliers = []
                for sheet_r, fv in row_val_map.items():
                    if fv < lower_bound or fv > upper_bound:
                        coord = f"{col_letter}{sheet_r}"
                        reason = f"Valor {fv} fuera de límites [{lower_bound:.2f}, {upper_bound:.2f}]"
                        col_outliers.append({
                            "cell": coord,
                            "row": sheet_r,
                            "value": fv,
                            "reason": reason
                        })
                        all_flagged_cells.append(coord)
                        outlier_cells_to_format.append((coord, sheet_r, abs_col_idx, fv, reason, add_comment, method))

                outliers_results.append({
                    "column": headers[target_col_rel - 1],
                    "method": method,
                    "total_values_evaluated": len(val_list),
                    "outliers_count": len(col_outliers),
                    "thresholds": thresholds,
                    "flagged_cells": [o["cell"] for o in col_outliers],
                    "outliers": col_outliers
                })

        # 5. Volcado a celdas OpenPyXL
        new_data_rows_count = len(data_rows)
        if clean_spec and clean_spec.get("trim_headers", True):
            for c_idx in range(1, num_cols + 1):
                ws.cell(row=min_row, column=min_col + c_idx - 1).value = headers[c_idx - 1]

        for r_offset, r_data in enumerate(data_rows):
            curr_r = min_row + 1 + r_offset
            for c_idx in range(1, num_cols + 1):
                cell = ws.cell(row=curr_r, column=min_col + c_idx - 1)
                val = r_data[c_idx - 1] if c_idx - 1 < len(r_data) else None
                cell.value = val
                if c_idx in col_date_cfg:
                    fmt = col_date_cfg[c_idx].get("format", "yyyy-mm-dd") if isinstance(col_date_cfg[c_idx], dict) else "yyyy-mm-dd"
                    cell.number_format = fmt

        # Eliminar filas sobrantes si hubo duplicados
        if duplicates_count > 0:
            leftover_start = min_row + 1 + new_data_rows_count
            ws.delete_rows(leftover_start, duplicates_count)
            if tbl:
                new_max_row = max_row - duplicates_count
                tbl.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{new_max_row}"

        # Formato visual de outliers
        if outlier_cells_to_format:
            warning_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
            warning_font = Font(color="C00000", bold=True)
            from openpyxl.comments import Comment
            for coord, sheet_r, abs_col_idx, fv, reason, add_com, m_name in outlier_cells_to_format:
                cell = ws[coord]
                cell.fill = warning_fill
                cell.font = warning_font
                if add_com:
                    try:
                        cell.comment = Comment(f"Outlier detectado ({m_name.upper()}): {reason}", "Antigravity")
                    except Exception:
                        pass

        return {
            "status": "success",
            "sheet": ws.title,
            "target_range_or_table": target_range_or_table or (tbl.name if tbl else f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"),
            "original_rows": orig_data_rows_count,
            "remaining_rows": len(data_rows),
            "duplicates_removed": duplicates_count,
            "cells_cleaned": cells_modified,
            "rules_applied": list(clean_spec.keys()) if clean_spec else [],
            "outliers_count": len(all_flagged_cells),
            "outliers_flagged": all_flagged_cells,
            "outliers_summary": outliers_results
        }

    def close(self):
        if self.wb:
            self.wb.close()

    def wait_until_ready(self, timeout_sec: float = 10.0, retry_interval: float = 0.25) -> Dict[str, Any]:
        """En modo headless (openpyxl), el motor siempre está inmediatamente listo."""
        return {
            "ready": True,
            "backend": "headless",
            "elapsed_sec": 0.0
        }

    def dispatch_com(
        self,
        target: str,
        member: str,
        args: Optional[List[Any]] = None,
        kwargs: Optional[Dict[str, Any]] = None
    ) -> Dict[str, Any]:
        """El despacho COM interactivo solo está disponible en el backend Live COM."""
        return {
            "success": False,
            "error": "dispatch_com solo está disponible en Live COM backend",
            "backend": "headless"
        }


    def open_workbook(self, file_path: str) -> Any:
        """Abre un libro de trabajo en memoria usando openpyxl."""
        self.file_path = file_path
        self.wb = self._load_workbook_resilient(file_path)
        return self.wb

    def get_sheet(self, sheet_name: Optional[str] = None) -> Any:
        """Obtiene la hoja de trabajo activa o la especificada por nombre."""
        return self._get_sheet(sheet_name)

    def supports_formula_recalc(self) -> bool:
        """Indica si el backend soporta motor nativo de recálculo dinámico de fórmulas."""
        return False

    def is_busy(self) -> bool:
        """Indica si el backend se encuentra ocupado o bloqueado."""
        return False

